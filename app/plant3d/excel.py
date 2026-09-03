from __future__ import annotations

import shutil
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0B4F6C"
TEAL = "0D7377"
HEADER_BG = "073B4C"
ROW_ALT = "E8F4F8"
META_LABEL = "5A6A75"
THIN = Border(
    left=Side(style="thin", color="C5D5DE"),
    right=Side(style="thin", color="C5D5DE"),
    top=Side(style="thin", color="C5D5DE"),
    bottom=Side(style="thin", color="C5D5DE"),
)

# English issued headers for common Plant 3D / Vitens column keys and Dutch labels.
ENGLISH_HEADER_BY_KEY: dict[str, str] = {
    "Tag": "Tag",
    "ObjectType": "Object type",
    "Omschrijving": "Description",
    "Size": "Size",
    "NONC": "NO/NC",
    "Actuation": "Actuation",
    "LineNumber": "Line number",
    "Service": "Service",
    "Medium": "Medium",
    "Material": "Material",
    "Procesdeel": "Process part",
    "Procesmodule": "Process module",
    "PnID": "Drawing",
    "Area": "Area",
    "Remarks": "Remarks",
    "PnId": "Drawing",
    "Description": "Description",
    "Type": "Type",
    "Class": "Class",
    "Status": "Status",
    "Spec": "Spec",
    "Rating": "Rating",
    "EndType": "End type",
    "Schedule": "Schedule",
    "Insulation": "Insulation",
    "Tracing": "Tracing",
    "Fluid": "Fluid",
    "DesignPressure": "Design pressure",
    "DesignTemperature": "Design temperature",
    "OperPressure": "Operating pressure",
    "OperTemperature": "Operating temperature",
    "Length": "Length",
    "Diameter": "Diameter",
    "NominalDiameter": "Nominal diameter",
    "ItemCode": "Item code",
    "Manufacturer": "Manufacturer",
    "Model": "Model",
    "SerialNumber": "Serial number",
}

ENGLISH_HEADER_BY_LABEL: dict[str, str] = {
    "Objectsoort": "Object type",
    "Omschrijving": "Description",
    "Maat": "Size",
    "Bediening": "Actuation",
    "Leidingnr": "Line number",
    "Materiaal": "Material",
    "Procesopstal": "Area",
    "Opmerking": "Remarks",
    "Blad": "Drawing",
    "Component": "Component",
    "Tekening": "Drawing",
    "Equipment": "Equipment",
    "Instrument": "Instrument",
    "Lijn": "Line",
    "Diameter": "Diameter",
    "Lengte": "Length",
}

LOGO_COL_END = 3
META_LABEL_COL = 4
META_VALUE_COL_START = 5
META_VALUE_COL_END = 7
REV_COL_START = 8
REV_COL_END = 13
TITLE_BLOCK_ROWS = 7
DATA_HEADER_ROW = 9


def english_column_header(col: dict[str, Any]) -> str:
    if col.get("header_en"):
        return str(col["header_en"])
    key = str(col.get("key") or "")
    header = str(col.get("header") or key)
    return ENGLISH_HEADER_BY_KEY.get(key) or ENGLISH_HEADER_BY_LABEL.get(header) or header


def _format_issue_date(value: Any) -> str:
    if not value:
        return date.today().isoformat()
    text = str(value).strip()
    if not text:
        return date.today().isoformat()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return text[:10]


def _style_cell(
    cell,
    *,
    bold: bool = False,
    size: int = 10,
    color: str = HEADER_BG,
    fill: str | None = None,
    align: str = "left",
    wrap: bool = False,
) -> None:
    cell.font = Font(name="Calibri", bold=bold, size=size, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _set_merged_value(ws, row: int, col_start: int, col_end: int, value: Any) -> None:
    if col_end > col_start:
        ws.merge_cells(
            start_row=row,
            start_column=col_start,
            end_row=row,
            end_column=col_end,
        )
    cell = ws.cell(row=row, column=col_start, value=value)
    return cell


def _build_title_block_meta(header: dict[str, Any], project: dict[str, Any], doc_no: str, rev: str) -> list[tuple[str, str]]:
    meta: list[tuple[str, str]] = []
    fields = header.get("fields") or []
    if fields:
        for field in fields:
            label = str(field.get("label") or field.get("key") or "").strip()
            value = str(field.get("value") or "").strip()
            if label:
                meta.append((label, value))
    else:
        meta.extend(
            [
                ("Project", str(project.get("name") or "")),
                ("Description", str(project.get("description") or "")),
                ("Status", str(project.get("status") or "")),
                ("Standard", str(project.get("standard") or "")),
            ]
        )
    meta.extend(
        [
            ("Document", doc_no),
            ("Revision", rev),
            ("Date", _format_issue_date(header.get("date"))),
        ]
    )
    return meta


def _write_title_block(
    ws,
    *,
    template: dict[str, Any],
    project: dict[str, Any],
    logo_path: Path | None,
    layout_cols: int,
) -> None:
    header = template.get("header") or {}
    revisions = template.get("revision_table") or []
    title = header.get("title") or template.get("name") or "Plant 3D List"
    doc_no = header.get("document_number") or f"{project.get('number', '')}-{template.get('id', 'LST')}"
    rev = header.get("revision") or (revisions[0].get("rev") if revisions else "A")
    meta = _build_title_block_meta(header, project, doc_no, rev)

    ws.merge_cells(start_row=1, start_column=1, end_row=TITLE_BLOCK_ROWS, end_column=LOGO_COL_END)
    logo_cell = ws.cell(1, 1)
    logo_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if logo_path and logo_path.exists():
        img = XLImage(str(logo_path))
        img.width = 132
        img.height = 46
        ws.add_image(img, "A1")
    else:
        logo_cell.value = header.get("company") or "COMPANY LOGO"
        _style_cell(logo_cell, bold=True, size=12, color=NAVY, align="center", wrap=True)

    _set_merged_value(ws, 1, META_LABEL_COL, META_VALUE_COL_END, title)
    title_cell = ws.cell(1, META_LABEL_COL)
    _style_cell(title_cell, bold=True, size=16, color=HEADER_BG, align="left")

    meta_start_row = 3
    for offset, (label, value) in enumerate(meta):
        row = meta_start_row + offset
        if row > TITLE_BLOCK_ROWS:
            break
        label_cell = ws.cell(row, META_LABEL_COL, label)
        _style_cell(label_cell, bold=True, size=9, color=META_LABEL)
        value_cell = _set_merged_value(ws, row, META_VALUE_COL_START, META_VALUE_COL_END, value)
        _style_cell(value_cell, size=10, color=HEADER_BG)

    rev_title = _set_merged_value(ws, 1, REV_COL_START, REV_COL_END, "REVISION HISTORY")
    _style_cell(rev_title, bold=True, size=9, color="FFFFFF", fill=NAVY, align="center")

    rev_headers = ["Rev", "Date", "Description", "Drawn", "Checked", "Approved"]
    for idx, name in enumerate(rev_headers):
        col = REV_COL_START + idx
        cell = ws.cell(2, col, name)
        _style_cell(cell, bold=True, size=8, color="FFFFFF", fill=TEAL, align="center")
        cell.border = THIN

    for r_i, rev_row in enumerate(revisions[:5], start=3):
        values = [
            rev_row.get("rev", ""),
            _format_issue_date(rev_row.get("date")),
            rev_row.get("desc", ""),
            rev_row.get("drawn", ""),
            rev_row.get("checked", ""),
            rev_row.get("approved", ""),
        ]
        for c_i, value in enumerate(values):
            col = REV_COL_START + c_i
            cell = ws.cell(r_i, col, value)
            _style_cell(cell, size=8, color=HEADER_BG, align="center" if c_i == 0 else "left")
            cell.border = THIN

    ws.row_dimensions[1].height = 24
    for row_idx in range(2, TITLE_BLOCK_ROWS + 1):
        ws.row_dimensions[row_idx].height = 16

    for col_idx in range(1, layout_cols + 1):
        letter = get_column_letter(col_idx)
        if col_idx <= LOGO_COL_END:
            ws.column_dimensions[letter].width = 14
        elif col_idx <= META_VALUE_COL_END:
            ws.column_dimensions[letter].width = 16 if col_idx == META_LABEL_COL else 18
        elif col_idx <= REV_COL_END:
            ws.column_dimensions[letter].width = 14 if col_idx == REV_COL_START + 2 else 11


def export_workbook(
    rows: list[dict[str, Any]],
    template: dict[str, Any],
    project: dict[str, Any],
    logo_path: Path | None = None,
) -> bytes:
    columns = [c for c in template.get("columns", []) if c.get("visible", True)]
    headers = [english_column_header(c) for c in columns]
    keys = [c["key"] for c in columns]
    include_id = template.get("include_pnpid", True)

    wb = Workbook()
    ws = wb.active
    ws.title = (template.get("name") or "List")[:31]

    header = template.get("header") or {}
    revisions = template.get("revision_table") or []
    doc_no = header.get("document_number") or f"{project.get('number', '')}-{template.get('id', 'LST')}"
    rev = header.get("revision") or (revisions[0].get("rev") if revisions else "A")

    data_col_count = len(keys) + (1 if include_id else 0)
    layout_cols = max(REV_COL_END, data_col_count)

    _write_title_block(ws, template=template, project=project, logo_path=logo_path, layout_cols=layout_cols)

    header_row = DATA_HEADER_ROW
    fill = PatternFill("solid", fgColor=HEADER_BG)
    font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    for idx, header_name in enumerate(headers, start=1):
        cell = ws.cell(header_row, idx, header_name)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN
    if include_id:
        cell = ws.cell(header_row, len(headers) + 1, "PnPID")
        cell.fill = fill
        cell.font = font
        cell.border = THIN

    alt = PatternFill("solid", fgColor=ROW_ALT)
    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, key in enumerate(keys, start=1):
            cell = ws.cell(r_idx, c_idx, row.get(key, ""))
            cell.font = Font(name="Calibri", size=9)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if (r_idx - header_row) % 2 == 0:
                cell.fill = alt
        if include_id:
            cell = ws.cell(r_idx, len(headers) + 1, row.get("PnPID", ""))
            cell.font = Font(name="Calibri", size=8, color="8A9BA8")
            cell.border = THIN

    last_data_col = len(keys) + (1 if include_id else 0)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(last_data_col)}{header_row + max(len(rows), 1)}"
    )
    ws.freeze_panes = f"A{header_row + 1}"
    ws.row_dimensions[header_row].height = 28

    for idx, col in enumerate(columns, start=1):
        width = col.get("width") or max(12, min(36, len(english_column_header(col)) + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width
    if include_id:
        ws.column_dimensions[get_column_letter(len(keys) + 1)].width = 10

    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.oddFooter.left.text = f"&8{doc_no}  Rev {rev}"
    ws.oddFooter.right.text = "&8Page &P of &N"

    meta_sheet = wb.create_sheet("_EasyReportCreator")
    meta_sheet["A1"] = "template_id"
    meta_sheet["B1"] = template.get("id", "")
    meta_sheet["A2"] = "source"
    meta_sheet["B2"] = template.get("source", "")
    meta_sheet["A3"] = "project"
    meta_sheet["B3"] = project.get("name", "")
    meta_sheet["A4"] = "export_language"
    meta_sheet["B4"] = "en"
    meta_sheet.sheet_state = "hidden"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_imported_excel(content: bytes) -> tuple[list[dict[str, Any]], dict[str, str]]:
    wb = load_workbook(BytesIO(content), data_only=True)
    meta = {}
    meta_name = next((n for n in ("_EasyReportCreator", "_PlantList") if n in wb.sheetnames), None)
    if meta_name:
        ms = wb[meta_name]
        for row in ms.iter_rows(min_row=1, max_col=2, values_only=True):
            if row[0]:
                meta[str(row[0])] = str(row[1] or "")
    ws = wb[wb.sheetnames[0]]
    header_row = 1
    for r in range(1, 20):
        values = [ws.cell(r, c).value for c in range(1, 12)]
        if any(v == "PnPID" for v in values) or (values[0] and str(values[0]).lower() in {"tag", "pnid"}):
            header_row = r
            break
    headers = []
    col = 1
    while True:
        value = ws.cell(header_row, col).value
        if value is None and col > 1:
            break
        headers.append(str(value) if value is not None else f"col{col}")
        col += 1
        if col > 80:
            break
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        item = {}
        empty = True
        for c, header in enumerate(headers, start=1):
            value = ws.cell(r, c).value
            if value not in (None, ""):
                empty = False
            item[header] = "" if value is None else value
        if not empty:
            rows.append(item)
    return rows, meta


def apply_changes_to_copy(
    source_dcf: Path,
    target_dcf: Path,
    changes: list[dict[str, Any]],
) -> Path:
    shutil.copy2(source_dcf, target_dcf)
    import sqlite3

    con = sqlite3.connect(target_dcf)
    try:
        for change in changes:
            con.execute(
                f"UPDATE EngineeringItems SET [{change['column']}] = ? WHERE PnPID = ?",
                (change["new"], change["pnpid"]),
            )
        con.commit()
    finally:
        con.close()
    return target_dcf
