"""WP6 Thu 3 Sep — control valves, instruments, drawings, line summary."""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from plant3d.dcf import connect, load_project
from plant3d.excel import english_column_header, export_workbook
from plant3d.project import merge_template_header
from plant3d.queries import run_source
from plant3d.templates import apply_template, load_template, resolve_template_id

SAMPLE = ROOT.parent / "samples" / "MN-P-RHN-PID-0001"

CHECKS = [
    ("control_valve_list", "control_valves", "control_valves", "Tag"),
    ("instrument_list", "instruments", "instruments", "Tag"),
    ("drawing_list", "drawings", "drawings", "Drawing"),
    ("line_summary", "line_summary", "line_groups", "Line number"),
]


def main() -> None:
    info, dcf_path = load_project(SAMPLE)
    print("Project:", info.name or info.number)
    print("Counts:", {k: info.counts[k] for k in ("control_valves", "instruments", "drawings", "line_groups")})

    with connect(dcf_path) as con:
        for template_id, source, count_key, expect_header in CHECKS:
            resolved = resolve_template_id(template_id)
            template = load_template(resolved)
            details = {"values": {}, "revisions": template.get("revision_table") or []}
            merged = merge_template_header(template, details)
            raw = run_source(con, source)
            rows = apply_template(raw, merged)
            expected = info.counts[count_key]
            assert len(raw) == expected, f"{source}: raw {len(raw)} != {expected}"
            assert len(rows) == expected, f"{resolved}: projected {len(rows)} != {expected}"

            if source == "line_summary":
                filled_ln = sum(1 for r in raw if str(r.get("LineNumber") or "").strip())
                assert filled_ln >= 300, f"line_summary LineNumber fill too low: {filled_ln}"
                filled_size = sum(1 for r in raw if str(r.get("Size") or "").strip())
                assert filled_size >= 300, f"line_summary Size fill too low: {filled_size}"

            workbook = export_workbook(rows, merged, info.__dict__, None)
            assert len(workbook) > 2000
            wb = load_workbook(BytesIO(workbook), data_only=True)
            ws = wb[wb.sheetnames[0]]
            headers = [ws.cell(9, c).value for c in range(1, 12) if ws.cell(9, c).value]
            assert expect_header in headers, f"{resolved}: missing English header {expect_header!r} in {headers}"

            for col in [c for c in merged["columns"] if c.get("visible", True)]:
                assert english_column_header(col)

            print(
                f"OK {source}: {len(rows)} rows, {len(workbook)} bytes xlsx, "
                f"headers={headers[:4]}…"
            )

    print("OK WP6 Thu lists: control valves, instruments, drawings, line summary")


if __name__ == "__main__":
    main()
