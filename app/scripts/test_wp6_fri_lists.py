"""WP6 Fri 4 Sep — Componentenlijst + regression on all eight list types."""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from plant3d.dcf import connect, load_project
from plant3d.excel import english_column_header, export_workbook
from plant3d.project import merge_template_header
from plant3d.queries import run_source
from plant3d.templates import apply_template, load_template, resolve_template_id

SAMPLE = ROOT.parent / "samples" / "MN-P-RHN-PID-0001"

# template_id, source, count_key, expected English header in Excel
CHECKS = [
    ("valve_list", "valves", "hand_valves", "Tag"),
    ("equipment_list", "equipment", "equipment", "Tag"),
    ("line_list", "lines", "pipe_lines", "Tag"),
    ("control_valve_list", "control_valves", "control_valves", "Tag"),
    ("instrument_list", "instruments", "instruments", "Tag"),
    ("drawing_list", "drawings", "drawings", "Drawing"),
    ("line_summary", "line_summary", "line_groups", "Line number"),
    ("component_list", "components", "components", "Old tag"),
]


def main() -> None:
    info, dcf_path = load_project(SAMPLE)
    print("Project:", info.name or info.number)
    print("Counts:", {k: info.counts.get(k) for _, _, k, _ in CHECKS})

    with connect(dcf_path) as con:
        for template_id, source, count_key, expect_header in CHECKS:
            resolved = resolve_template_id(template_id)
            template = load_template(resolved)
            details = {"values": {}, "revisions": template.get("revision_table") or []}
            merged = merge_template_header(template, details)
            raw = run_source(con, source)
            rows = apply_template(raw, merged)
            expected = info.counts[count_key]
            assert len(raw) == expected, f"{source}: raw {len(raw)} != count {expected}"
            assert len(rows) == expected, f"{resolved}: projected {len(rows)} != {expected}"

            if source == "components":
                assert expected >= 800, f"components count unexpectedly low: {expected}"
                tagged = sum(1 for r in raw if str(r.get("Tag") or "").strip())
                assert tagged >= 500, f"components tagged rows too low: {tagged}"

            workbook = export_workbook(rows, merged, info.__dict__, None)
            assert len(workbook) > 2000
            wb = load_workbook(BytesIO(workbook), data_only=True)
            ws = wb[wb.sheetnames[0]]
            headers = [ws.cell(9, c).value for c in range(1, 16) if ws.cell(9, c).value]
            assert expect_header in headers, (
                f"{resolved}: missing English header {expect_header!r} in {headers}"
            )
            for col in [c for c in merged["columns"] if c.get("visible", True)]:
                assert english_column_header(col)

            print(
                f"OK {source}: {len(rows)} rows, {len(workbook)} bytes xlsx, "
                f"headers={headers[:4]}…"
            )

    print("OK WP6 Fri: Componentenlijst + all eight list types")


if __name__ == "__main__":
    main()
