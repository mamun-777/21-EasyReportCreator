"""WP5 Excel export checks (Tue 1 Sep)."""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from plant3d.dcf import connect, load_project
from plant3d.excel import export_workbook, english_column_header
from plant3d.project import merge_template_header
from plant3d.queries import run_source
from plant3d.templates import apply_template, load_template, resolve_template_id

SAMPLE = ROOT.parent / "samples" / "MN-P-RHN-PID-0001"


def main() -> None:
    info, dcf_path = load_project(SAMPLE)
    template_id = resolve_template_id("valve_list")
    template = load_template(template_id)
    with connect(dcf_path) as con:
        details = {"values": {}, "revisions": template.get("revision_table", [])}
        merged = merge_template_header(template, details)
        raw = run_source(con, merged["source"])
    rows = apply_template(raw, merged)
    workbook = export_workbook(rows, merged, info.__dict__, None)
    assert len(workbook) > 5000

    wb = load_workbook(BytesIO(workbook), data_only=True)
    ws = wb[wb.sheetnames[0]]

    assert ws["D1"].value and "Valve" in str(ws["D1"].value)
    assert ws.cell(9, 1).value == "Tag"
    assert ws.cell(9, 2).value == "Object type"
    assert ws.cell(9, 4).value == "Line number"
    assert ws["H1"].value == "REVISION HISTORY"
    assert ws["H3"].value == "B"

    visible = [c for c in merged["columns"] if c.get("visible", True)]
    for col in visible:
        assert english_column_header(col)

    print("OK excel export:", len(rows), "rows,", len(workbook), "bytes,", len(visible), "columns")


if __name__ == "__main__":
    main()
